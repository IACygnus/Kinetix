"""ETAPA H8.5b — el importador de proyectos en pantalla (H-D94, H-D99, H-D101).

    docker exec -e KX_API_PUERTO=8002 jmeter_backend python3 /tmp/e2e/h85b_pantalla.py

Lo que se comprueba aquí y no en el backend:

- que hay **dos pestañas** y que la de registros sigue funcionando (H-D94: el
  importador de horas no se toca);
- que la **plantilla se descarga** desde la pantalla (H-D101);
- que la previa enseña **un bloque por proyecto** con su total, y las filas malas
  con su motivo;
- y que **hasta confirmar no se ha escrito nada**.

Datos `ZZTEST-H85BP` (regla 29), base de pruebas (regla 34). 0 llamadas a la IA.
"""
import io
import json
import os
import subprocess
import sys

import httpx
from playwright.sync_api import sync_playwright

WEB = os.environ.get("KX_WEB", "http://localhost:5173")
API = os.environ.get("KX_API", "http://localhost:8002/api/v1")
BASE = os.environ.get("KX_DB", "jmeter_analyzer_test")
SESION = os.environ.get("KX_SESION", "/tmp/e2e_sesion_test.json")
PUERTO_API = os.environ.get("KX_API_PUERTO", "8002")

MARCA = "ZZTEST-H85BP"
CLIENTE = f"{MARCA}-cliente"
LIBRO = "/tmp/zztest_h85bp.xlsx"
DESCARGAS = "/tmp/descargas_h85bp"

fallos = []


def ok(cond, texto):
    print(f"{'PASA ' if cond else 'FALLA'} | {texto}")
    if not cond:
        fallos.append(texto)
    return cond


def desviar(page):
    if PUERTO_API == "8001":
        return
    page.route("http://localhost:8001/**", lambda r: r.continue_(
        url=r.request.url.replace("localhost:8001", f"localhost:{PUERTO_API}")))


def psql(sql):
    return subprocess.run(
        ["psql", "-U", "jmeter_user", "-d", BASE, "-q", "-t", "-A", "-c", sql],
        env={**os.environ, "PGHOST": "postgres", "PGPASSWORD": "jmeter_secure_2024"},
        capture_output=True, text=True)


def limpiar():
    if "test" not in BASE:
        sys.exit(f"PARADA: «{BASE}» no es una base de pruebas. No se borra nada.")
    p = f"name like '{MARCA}%'"
    psql(
        f"delete from time_entries where project_id in (select id from projects where {p});"
        f"delete from project_status_changes where project_id in (select id from projects where {p});"
        f"delete from project_activity_changes where project_id in (select id from projects where {p});"
        f"delete from project_activities where project_id in (select id from projects where {p});"
        f"delete from projects where {p};"
        f"delete from user_clients where client_id in (select id from clients where {p});"
        f"delete from clients where {p};"
        f"delete from activities where name like '{MARCA}%';"
    )


def escribir_libro():
    from openpyxl import Workbook
    wb = Workbook()
    hoja = wb.active
    hoja.append(["Cliente", "Proyecto", "Estado", "Actividad", "Horas estimadas"])
    hoja.append([CLIENTE, f"{MARCA}-alfa", "Pendiente", f"{MARCA}-planeacion", 16])
    hoja.append([CLIENTE, f"{MARCA}-alfa", "Pendiente", f"{MARCA}-ejecucion", 24])
    hoja.append([CLIENTE, f"{MARCA}-beta", "Detenido", f"{MARCA}-planeacion", 8])
    # Una mala, para ver el motivo en pantalla (H-D100).
    hoja.append([CLIENTE, f"{MARCA}-malo", "archivado", f"{MARCA}-planeacion", 8])
    wb.save(LIBRO)


def cuenta_proyectos():
    return int(psql(f"select count(*) from projects where name like '{MARCA}%';")
               .stdout.strip() or 0)


def main():
    limpiar()
    escribir_libro()
    os.makedirs(DESCARGAS, exist_ok=True)

    ck = {c["name"]: c["value"] for c in json.load(open(SESION))["cookies"]}
    cli = httpx.Client(cookies=ck, headers={"X-CSRF-Token": ck.get("csrf_token", "")},
                       timeout=120.0)
    if cli.get(f"{API}/auth/me").status_code != 200:
        sys.exit("sesion caducada: corre refrescar_sesion.py")

    with sync_playwright() as p:
        nav = p.chromium.launch(args=["--no-sandbox", "--disable-dev-shm-usage"])
        ctx = nav.new_context(viewport={"width": 1600, "height": 1500},
                              storage_state=SESION, locale="es-CO",
                              accept_downloads=True)
        page = ctx.new_page()
        desviar(page)
        errores = []
        page.on("pageerror", lambda e: errores.append(str(e)))

        print("=== 1. Las dos pestañas (H-D94) ===")
        page.goto(f"{WEB}/horas/importar", wait_until="networkidle")
        page.wait_for_selector("[data-testid='pestanas-importar']", timeout=25000)
        ok(page.locator("[data-testid='pestana-registros']").count() == 1,
           "está «Registros de horas»")
        ok(page.locator("[data-testid='pestana-proyectos']").count() == 1,
           "y «Proyectos y estimaciones»")
        ok(page.locator("[data-testid='archivo']").count() == 1,
           "y de entrada se ve el importador de horas, como antes")
        ok(page.locator("[data-testid='importar-proyectos']").count() == 0,
           "el nuevo está en su pestaña, sin estorbar")

        page.locator("[data-testid='pestana-proyectos']").click()
        page.wait_for_selector("[data-testid='importar-proyectos']", timeout=10000)
        ok(page.locator("[data-testid='archivo']").count() == 0,
           "al cambiar de pestaña, el de horas se retira")

        print("\n=== 2. La plantilla (H-D101) ===")
        with page.expect_download(timeout=20000) as dl:
            page.locator("[data-testid='descargar-plantilla']").click()
        descarga = dl.value
        ruta = os.path.join(DESCARGAS, descarga.suggested_filename)
        descarga.save_as(ruta)
        ok(descarga.suggested_filename.endswith(".xlsx"),
           f"se descarga: {descarga.suggested_filename}")
        from openpyxl import load_workbook
        wb = load_workbook(ruta)
        titulos = [c.value for c in wb.worksheets[0][1]]
        ok(titulos == ["Cliente", "Proyecto", "Estado", "Actividad", "Horas estimadas"],
           f"con las cinco columnas: {titulos}")

        print("\n=== 3. La vista previa ===")
        antes = cuenta_proyectos()
        page.locator("[data-testid='archivo-proyectos']").set_input_files(LIBRO)
        page.locator("[data-testid='analizar-proyectos']").click()
        page.wait_for_selector("[data-testid='previa-proyectos']", timeout=25000)

        bloques = page.locator("[data-testid='bloque-proyecto']")
        ok(bloques.count() == 2, f"dos bloques, uno por proyecto: {bloques.count()}")
        alfa = page.locator(f"[data-proyecto='{MARCA}-alfa']")
        ok(alfa.count() == 1, "el proyecto con dos actividades está")
        total = alfa.locator("[data-testid='total-proyecto']").inner_text().strip()
        ok(total.startswith("40"),
           f"y dice lo que suma el proyecto entero: «{total}» (16 + 24)")
        ok(alfa.locator("tbody tr").count() == 2,
           "con sus dos actividades juntas (H-D96)")
        ok(alfa.locator("[data-testid='chip-estado']").inner_text().strip() == "Pendiente",
           "y su estado del archivo")

        malas = page.locator("[data-testid='filas-invalidas']")
        ok(malas.count() == 1, "las filas que no entran salen aparte")
        ok("no es un estado" in malas.inner_text(),
           f"con su motivo: «{' '.join(malas.inner_text().split())[:80]}»")

        ok(cuenta_proyectos() == antes,
           "**y hasta aquí no se ha escrito nada** (H-D99)")

        print("\n=== 4. La importación ===")
        page.locator("[data-testid='confirmar-proyectos']").click()
        page.wait_for_selector("[data-testid='resumen-proyectos']", timeout=25000)
        texto = " ".join(page.locator("[data-testid='resumen-proyectos']").inner_text().split())
        ok("2 proyectos creados" in texto, f"el resumen: «{texto[:110]}»")
        ok("1 filas no se importaron" in texto or "1 fila" in texto,
           "y dice cuántas se quedaron fuera")
        ok(cuenta_proyectos() == antes + 2,
           f"y ahora sí están en la base: {cuenta_proyectos()}")

        estados = dict(
            linea.split("|") for linea in psql(
                f"select name, status from projects where name like '{MARCA}%';"
            ).stdout.strip().splitlines() if "|" in linea)
        ok(estados.get(f"{MARCA}-alfa") == "pendiente",
           f"«alfa» quedó en «{estados.get(f'{MARCA}-alfa')}»")
        ok(estados.get(f"{MARCA}-beta") == "detenido",
           f"«beta» quedó en «{estados.get(f'{MARCA}-beta')}»")

        print("\n=== 5. La pestaña de horas sigue entera (H-D94) ===")
        page.locator("[data-testid='pestana-registros']").click()
        page.wait_for_selector("[data-testid='archivo']", timeout=10000)
        for t in ("archivo", "importar-persona", "analizar"):
            ok(page.locator(f"[data-testid='{t}']").count() == 1,
               f"sigue el control «{t}»")

        ok(not errores, f"la consola queda limpia ({len(errores)}): {errores[:2]}")
        ctx.close()
        nav.close()

    print()
    if fallos:
        print(f"=== {len(fallos)} FALLO(S) ===")
        for f in fallos:
            print(f"  - {f}")
        return 1
    print("=== TODO PASA ===")
    return 0


if __name__ == "__main__":
    codigo = 1
    try:
        codigo = main()
    finally:
        limpiar()
    sys.exit(codigo)
