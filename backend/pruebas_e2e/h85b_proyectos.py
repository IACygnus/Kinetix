"""ETAPA H8.5b — importar proyectos, estados y estimaciones (H-D94 a H-D107).

    docker exec jmeter_backend python3 /tmp/e2e/h85b_proyectos.py

El archivo de prueba es el que pidió Fredy, todo en uno:

  - un proyecto NUEVO con tres actividades (H-D96);
  - un proyecto que YA EXISTE, para que actualice (H-D97);
  - **un estado de cada uno de los cinco** (H-D95);
  - una fila con un estado inventado (H-D100);
  - una con las horas mal (H-D100);
  - y una sin actividad (H-D100).

Y después se **reimporta el mismo archivo** para comprobar que las horas
estimadas no cambian, que es lo que significa idempotente.

Datos `ZZTEST-H85B` (regla 29), base de pruebas (regla 34). 0 llamadas a la IA.
"""
import io
import json
import os
import subprocess
import sys
from decimal import Decimal

import httpx

API = os.environ.get("KX_API", "http://localhost:8002/api/v1")
BASE = os.environ.get("KX_DB", "jmeter_analyzer_test")
SESION = os.environ.get("KX_SESION", "/tmp/e2e_sesion_test.json")

MARCA = "ZZTEST-H85B"
CLIENTE = f"{MARCA}-cliente"
NUEVO = f"{MARCA}-nuevo"
EXISTE = f"{MARCA}-existe"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

fallos = []


def ok(cond, texto):
    print(f"{'PASA ' if cond else 'FALLA'} | {texto}")
    if not cond:
        fallos.append(texto)
    return cond


def psql(sql):
    return subprocess.run(
        ["psql", "-U", "jmeter_user", "-d", BASE, "-q", "-t", "-A", "-c", sql],
        env={**os.environ, "PGHOST": "postgres", "PGPASSWORD": "jmeter_secure_2024"},
        capture_output=True, text=True)


def limpiar():
    if "test" not in BASE:
        sys.exit(f"PARADA: «{BASE}» no es una base de pruebas. No se borra nada.")
    p = f"name like '{MARCA}%'"
    psql(
        f"delete from time_entries where project_id in (select id from projects where {p});"
        f"delete from project_status_changes where project_id in (select id from projects where {p});"
        f"delete from project_activity_changes where project_id in (select id from projects where {p});"
        f"delete from project_activities where project_id in (select id from projects where {p});"
        f"delete from projects where {p};"
        f"delete from user_clients where client_id in (select id from clients where {p});"
        f"delete from clients where {p};"
        f"delete from activities where name like '{MARCA}%';"
    )


def libro(filas, titulos=None):
    from openpyxl import Workbook
    wb = Workbook()
    hoja = wb.active
    hoja.append(titulos or ["Cliente", "Proyecto", "Estado", "Actividad", "Horas estimadas"])
    for f in filas:
        hoja.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def enviar(cli, ruta, datos):
    return cli.post(f"{API}/time/import/proyectos/{ruta}",
                    files={"archivo": ("zztest_h85b.xlsx", datos, XLSX)})


def main():
    limpiar()
    ck = {c["name"]: c["value"] for c in json.load(open(SESION))["cookies"]}
    cli = httpx.Client(cookies=ck, headers={"X-CSRF-Token": ck.get("csrf_token", "")},
                       timeout=180.0)
    if cli.get(f"{API}/auth/me").status_code != 200:
        sys.exit("sesion caducada: corre refrescar_sesion.py")

    print("=== 0. La plantilla (H-D101) ===")
    r = cli.get(f"{API}/time/import/proyectos/plantilla")
    ok(r.status_code == 200 and r.content[:2] == b"PK",
       f"se descarga y es un .xlsx ({len(r.content)} bytes)")
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    titulos = [c.value for c in wb.worksheets[0][1]]
    ok(titulos == ["Cliente", "Proyecto", "Estado", "Actividad", "Horas estimadas"],
       f"con las cinco columnas: {titulos}")
    ejemplos = [f for f in wb.worksheets[0].iter_rows(min_row=2, values_only=True)]
    ok(len(ejemplos) == 2, f"y dos filas de ejemplo ({len(ejemplos)})")
    ok(ejemplos[0][1] == ejemplos[1][1],
       "del MISMO proyecto: así se ve que una fila es una actividad (H-D96)")
    ok("Cómo se llena" in wb.sheetnames,
       "la ayuda va en otra hoja, para que no se lea como una fila más")

    # El proyecto que YA existe, creado por la vía normal.
    existe = cli.post(f"{API}/time/projects", json={
        "client_id": cli.post(f"{API}/clients", json={"name": CLIENTE}).json()["id"],
        "name": EXISTE,
        "activities": [{"activity_id": cli.get(f"{API}/time/activities").json()[0]["id"],
                        "estimated_hours": 10}]})
    if existe.status_code != 201:
        sys.exit(f"no se pudo preparar el proyecto existente: {existe.text[:200]}")
    act_existente = cli.get(f"{API}/time/activities").json()[0]["name"]
    n_proyectos_antes = int(psql("select count(*) from projects;").stdout.strip())

    datos = libro([
        # Proyecto NUEVO con TRES actividades (H-D96), y un estado de cada uno.
        [CLIENTE, NUEVO, "Pendiente", f"{MARCA}-planeacion", 16],
        [CLIENTE, NUEVO, "Pendiente", f"{MARCA}-diseno", 24],
        [CLIENTE, NUEVO, "Pendiente", f"{MARCA}-ejecucion", 40],
        # El que ya existe: cambia de estado y ACTUALIZA la estimación (H-D97).
        [CLIENTE, EXISTE, "Detenido", act_existente, 18],
        # Los otros tres estados, cada uno con su proyecto.
        [CLIENTE, f"{MARCA}-en-ejecucion", "En ejecución", f"{MARCA}-planeacion", 8],
        [CLIENTE, f"{MARCA}-no-viable", "No viable", f"{MARCA}-planeacion", 4],
        [CLIENTE, f"{MARCA}-finalizado", "Finalizado", f"{MARCA}-planeacion", 4],
        # Las tres malas (H-D100).
        [CLIENTE, f"{MARCA}-malo", "archivado", f"{MARCA}-planeacion", 8],
        [CLIENTE, f"{MARCA}-malo", "Pendiente", f"{MARCA}-planeacion", "ocho"],
        [CLIENTE, f"{MARCA}-malo", "Pendiente", "", 8],
    ])

    print("\n=== 1. La vista previa ===")
    r = enviar(cli, "preview", datos)
    if not ok(r.status_code == 200, f"responde {r.status_code}: {r.text[:200]}"):
        return 1
    v = r.json()
    ok(v["total_filas"] == 10, f"lee las 10 filas ({v['total_filas']})")
    ok(len(v["invalidas"]) == 3, f"y aparta 3 (H-D100): {len(v['invalidas'])}")
    motivos = {f["numero"]: f["motivo"] for f in v["invalidas"]}
    ok(any("no es un estado" in m for m in motivos.values()),
       f"el estado inventado, con su motivo: {motivos}")
    ok(any("no se entienden" in m for m in motivos.values()), "las horas ilegibles")
    ok(any("Falta Actividad" in m for m in motivos.values()), "y la fila sin actividad")

    bloques = {b["project_name"]: b for b in v["proyectos"]}
    ok(len(bloques) == 5, f"quedan 5 proyectos ({len(bloques)}): {sorted(bloques)}")
    if ok(NUEVO in bloques, "el proyecto nuevo está"):
        b = bloques[NUEVO]
        ok(b["actividades"] == 3, f"con sus 3 actividades juntas (H-D96): {b['actividades']}")
        ok(Decimal(str(b["total_hours"])) == Decimal("80"),
           f"y lo que suma el proyecto entero: {b['total_hours']} h")
        ok(b["es_nuevo"] is True, "marcado como nuevo")
    if ok(EXISTE in bloques, "y el que ya existía"):
        b = bloques[EXISTE]
        ok(b["es_nuevo"] is False, "marcado como existente")
        ok(b["cambia_de_estado"] is True,
           f"y avisa de que cambia de estado: «{b['status_anterior_label']}» → "
           f"«{b['status_label']}»")
        ok(b["filas"][0]["accion"] == "actualiza",
           f"su estimación se actualiza: {b['filas'][0]['accion']}")
        ok(Decimal(str(b["filas"][0]["previous_hours"])) == Decimal("10"),
           f"y dice lo que había: {b['filas'][0]['previous_hours']} h")

    ok(v["estimaciones_nuevas"] == 6, f"6 estimaciones nuevas: {v['estimaciones_nuevas']}")
    ok(v["estimaciones_actualizadas"] == 1,
       f"1 actualizada: {v['estimaciones_actualizadas']}")
    ok(len(v["actividades_a_crear"]) == 3,
       f"y 3 actividades a crear (H-D98): {v['actividades_a_crear']}")

    # Se compara contra la foto de antes, no contra un número fijo: la base de
    # pruebas tiene proyectos de otras etapas que no son míos y no se tocan
    # (regla 30).
    ok(int(psql("select count(*) from projects;").stdout.strip()) == n_proyectos_antes,
       f"**y la previa no escribió nada**: siguen {n_proyectos_antes} proyecto(s)")

    print("\n=== 2. La importación ===")
    r = enviar(cli, "confirm", datos)
    if not ok(r.status_code == 200, f"responde {r.status_code}: {r.text[:300]}"):
        return 1
    res = r.json()
    ok(len(res["proyectos_creados"]) == 4,
       f"crea 4 proyectos: {[p['name'] for p in res['proyectos_creados']]}")
    ok(res["proyectos_actualizados"] == 1, f"y toca 1: {res['proyectos_actualizados']}")
    ok(res["estimaciones_creadas"] == 6, f"6 estimaciones nuevas: {res['estimaciones_creadas']}")
    ok(res["estimaciones_actualizadas"] == 1, f"1 actualizada: {res['estimaciones_actualizadas']}")
    ok(res["estados_cambiados"] == 1, f"1 cambio de estado: {res['estados_cambiados']}")
    ok(res["omitidas"] == 3, f"y 3 filas fuera: {res['omitidas']}")

    print("\n=== 3. Los cinco estados quedaron puestos (H-D95) ===")
    esperado = {NUEVO: "pendiente", EXISTE: "detenido",
                f"{MARCA}-en-ejecucion": "en_ejecucion",
                f"{MARCA}-no-viable": "no_viable",
                f"{MARCA}-finalizado": "finalizado"}
    for nombre, estado in esperado.items():
        real = psql(f"select status from projects where name = '{nombre}';").stdout.strip()
        ok(real == estado, f"«{nombre}» quedó en «{real}»")

    print("\n=== 4. Las estimaciones, y el historial ===")
    total = psql(f"select coalesce(sum(estimated_hours),0) from project_activities "
                 f"where project_id in (select id from projects where name like '{MARCA}%');")
    ok(Decimal(total.stdout.strip()) == Decimal("114.00"),
       f"suman 114 h (16+24+40+18+8+4+4): {total.stdout.strip()}")
    cambios = psql(f"select change_type from project_activity_changes where project_id in "
                   f"(select id from projects where name = '{EXISTE}') order by changed_at;")
    ok("cambio" in cambios.stdout,
       f"el cambio de estimación dejó su historial: {cambios.stdout.split()}")

    print("\n=== 5. Reimportar el MISMO archivo no cambia nada (H-D97) ===")
    r = enviar(cli, "confirm", datos)
    if ok(r.status_code == 200, f"responde {r.status_code}"):
        res2 = r.json()
        ok(len(res2["proyectos_creados"]) == 0,
           f"no crea ningún proyecto: {len(res2['proyectos_creados'])}")
        ok(res2["estimaciones_creadas"] == 0,
           f"ni ninguna estimación: {res2['estimaciones_creadas']}")
        ok(res2["estimaciones_actualizadas"] == 0,
           f"ni actualiza: {res2['estimaciones_actualizadas']}")
        ok(res2["estimaciones_iguales"] == 7,
           f"las 7 salen «iguales»: {res2['estimaciones_iguales']}")
        ok(res2["estados_cambiados"] == 0, "y ningún estado cambia")

    total2 = psql(f"select coalesce(sum(estimated_hours),0) from project_activities "
                  f"where project_id in (select id from projects where name like '{MARCA}%');")
    ok(total2.stdout.strip() == total.stdout.strip(),
       f"**las horas estimadas no cambian**: {total2.stdout.strip()}")
    n_hist = psql(f"select count(*) from project_activity_changes where project_id in "
                  f"(select id from projects where name like '{MARCA}%');").stdout.strip()
    # 8 = el «alta» del proyecto que ya existía (creado por la vía normal)
    # + las 6 altas de la importación + el «cambio» de la estimación que subió
    # de 10 a 18. Reimportar no añade ninguna más.
    ok(n_hist == "8", f"y no se escribió historial de más: {n_hist} líneas")

    print("\n=== 6. La columna «Estado» es opcional (H-D95) ===")
    sin_estado = libro(
        [[CLIENTE, f"{MARCA}-sin-estado", f"{MARCA}-planeacion", 8]],
        titulos=["Cliente", "Proyecto", "Actividad", "Horas estimadas"])
    r = enviar(cli, "preview", sin_estado)
    if ok(r.status_code == 200, f"un archivo SIN esa columna se acepta: {r.status_code}"):
        b = r.json()["proyectos"][0]
        ok(b["status"] == "en_ejecucion",
           f"y el proyecto queda «{b['status_label']}», que es el valor por defecto")

    print("\n=== 7. Un archivo sin una columna obligatoria ===")
    sin_cliente = libro([["x", "y", 8]],
                        titulos=["Proyecto", "Actividad", "Horas estimadas"])
    r = enviar(cli, "preview", sin_cliente)
    ok(r.status_code == 400 and "Cliente" in r.json().get("detail", ""),
       f"se para antes de leer filas y dice cuál falta: "
       f"{r.json().get('detail', '')[:80]}")

    print("\n=== 8. El orden de las columnas da igual (H-D95) ===")
    al_reves = libro(
        [[8, f"{MARCA}-planeacion", "Pendiente", f"{MARCA}-orden", CLIENTE]],
        titulos=["Horas estimadas", "Actividad", "Estado ", "Proyecto", "Cliente"])
    r = enviar(cli, "preview", al_reves)
    if ok(r.status_code == 200, f"con las columnas cambiadas de sitio: {r.status_code}"):
        b = r.json()["proyectos"][0]
        ok(b["project_name"] == f"{MARCA}-orden" and b["status"] == "pendiente",
           f"se leen igual, y «Estado » con espacio final también: {b['status_label']}")

    print("\n=== 9. La base de Fredy, intacta ===")
    suyos = subprocess.run(
        ["psql", "-U", "jmeter_user", "-d", "jmeter_analyzer_db", "-q", "-t", "-A",
         "-c", "select count(*) from projects;"],
        env={**os.environ, "PGHOST": "postgres", "PGPASSWORD": "jmeter_secure_2024"},
        capture_output=True, text=True).stdout.strip()
    # El número cambia cuando Fredy carga o borra de verdad: la CARGA REAL del
    # 24 de septiembre de 2026 dejó su base en 16 proyectos (eran 10). Se
    # actualiza cuando eso pasa **a sabiendas**; que falle sin que nadie haya
    # cargado nada es justo lo que esta comprobación tiene que delatar.
    ESPERADOS = os.environ.get("KX_PROYECTOS_FREDY", "16")
    ok(suyos == ESPERADOS,
       f"sigue con sus {ESPERADOS} proyectos: {suyos}")

    print()
    if fallos:
        print(f"=== {len(fallos)} FALLO(S) ===")
        for f in fallos:
            print(f"  - {f}")
        return 1
    print("=== TODO PASA ===")
    return 0


if __name__ == "__main__":
    codigo = 1
    try:
        codigo = main()
    finally:
        limpiar()
    sys.exit(codigo)
