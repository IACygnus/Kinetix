"""El lector del archivo de proyectos y estimaciones (ETAPA H8.5b, H-D94 a H-D101).

Segundo importador. **El de registros no se toca** (H-D94): de él se reutiliza lo
que ya está medido contra archivos reales —`filas_del_libro`, `leer_horas`,
`limpiar_texto`, `FaltanColumnas`— y aquí solo vive lo propio de este formato.

Misma separación que en H3, y por la misma razón: **interpretar un valor son
funciones puras** que se prueban sin fabricar un `.xlsx`, y **abrir el libro** es
lo único que necesita `openpyxl`.

Las cinco columnas (H-D95):

    Cliente · Proyecto · Estado · Actividad · Horas estimadas

Cliente, Proyecto, Actividad y Horas estimadas son obligatorias; **Estado es
opcional** y si falta vale `en_ejecucion`. Los títulos se comparan
**normalizados**, así que el orden da igual y los espacios al final tampoco
estorban.

**Una fila por actividad del proyecto** (H-D96): varias filas del mismo proyecto
son sus distintas actividades, no un proyecto repetido.
"""
import io
from decimal import Decimal
from typing import Any, Dict, List, Optional

from app.db.models.time_tracking import normalizar
from app.services.horas import estados
# Lo que ya está medido contra el archivo real de Fredy no se vuelve a escribir.
from app.services.horas.importacion import (  # noqa: F401  (FaltanColumnas se reexporta)
    FaltanColumnas, leer_horas, limpiar_texto,
)

# clave interna -> título tal como se lee en el archivo (H-D95).
COLUMNAS: Dict[str, str] = {
    "cliente": "Cliente",
    "proyecto": "Proyecto",
    "estado": "Estado",
    "actividad": "Actividad",
    "horas": "Horas estimadas",
}

#: «Estado» NO está: es opcional (H-D95).
OBLIGATORIAS = ("cliente", "proyecto", "actividad", "horas")


def mapear_columnas(titulos: List[Any]) -> Dict[str, int]:
    """De la fila de títulos a `{clave: posición}`, comparando normalizado.

    Si falta una obligatoria se para aquí, con el nombre de la que falta: leer
    doscientas filas para descubrirlo al final no ayuda a nadie.
    """
    vistos = {normalizar(t): i for i, t in enumerate(titulos)
              if t is not None and str(t).strip()}
    mapa = {clave: vistos[normalizar(titulo)]
            for clave, titulo in COLUMNAS.items()
            if normalizar(titulo) in vistos}
    faltan = [COLUMNAS[c] for c in OBLIGATORIAS if c not in mapa]
    if faltan:
        raise FaltanColumnas(faltan)
    return mapa


# Lo que se acepta escrito en la columna «Estado»: la clave interna o el rótulo
# que se lee en pantalla, los dos normalizados. Así el archivo puede traer
# «En ejecución», «en_ejecucion» o «EN EJECUCION» y significan lo mismo.
_ESTADOS_ACEPTADOS: Dict[str, str] = {}
for _e in estados.ESTADOS:
    _ESTADOS_ACEPTADOS[normalizar(_e)] = _e
    _ESTADOS_ACEPTADOS[normalizar(_e.replace("_", " "))] = _e
    _ESTADOS_ACEPTADOS[normalizar(estados.TEXTOS[_e])] = _e


def leer_estado(valor: Any) -> Optional[str]:
    """El estado de la fila (H-D95, H-D100).

    - vacío  → `en_ejecucion`, que es el valor por defecto de §3.1;
    - conocido → su clave interna;
    - **cualquier otra cosa → `None`**, y la fila se descarta con su motivo. No
      se adivina: «terminado» se parece a «finalizado» y no es lo mismo, y
      elegir por parecido metería un estado que nadie escribió.

    >>> leer_estado("En Ejecución")
    'en_ejecucion'
    >>> leer_estado("")
    'en_ejecucion'
    >>> leer_estado("archivado") is None
    True
    """
    texto = limpiar_texto(valor)
    if not texto:
        return estados.POR_DEFECTO
    return _ESTADOS_ACEPTADOS.get(normalizar(texto))


def estados_para_ayuda() -> str:
    """Los cinco, tal como se pueden escribir. Para el mensaje de una fila mala."""
    return " · ".join(estados.TEXTOS[e] for e in estados.ESTADOS)


# ===================== LA PLANTILLA (H-D101) =====================

def plantilla_xlsx() -> bytes:
    """Un `.xlsx` con las cinco columnas y dos filas de ejemplo.

    Las dos filas son **del mismo proyecto** a propósito: es la forma de enseñar
    sin explicar que una fila es una actividad y que el proyecto se repite
    (H-D96).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Proyectos"
    titulos = [COLUMNAS[c] for c in ("cliente", "proyecto", "estado",
                                     "actividad", "horas")]
    hoja.append(titulos)
    for celda in hoja[1]:
        celda.font = Font(bold=True)

    hoja.append(["Banco Ejemplo", "Migración core", "En ejecución",
                 "Planeación", 16])
    hoja.append(["Banco Ejemplo", "Migración core", "En ejecución",
                 "Ejecución", 40])

    anchos = (22, 28, 16, 30, 16)
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[hoja.cell(row=1, column=i).column_letter].width = ancho

    # La ayuda va en una SEGUNDA hoja, no debajo de la tabla: el lector toma la
    # primera hoja del libro (H-D47) y una nota al pie en la columna «Cliente»
    # se leería como una fila más, y saldría en la previa como inválida.
    ayuda = wb.create_sheet("Cómo se llena")
    for i, linea in enumerate((
        "Una fila por actividad del proyecto.",
        "El proyecto se repite en cada una de sus actividades: no se duplica, "
        "se le añaden todas.",
        "Cliente, Proyecto, Actividad y Horas estimadas son obligatorias.",
        f"Estado es opcional; si se deja vacío vale «{estados.TEXTOS[estados.POR_DEFECTO]}».",
        f"Estados que se aceptan: {estados_para_ayuda()}.",
        "Las horas estimadas van en pasos de 0,25 y tienen que ser mayores que cero.",
        "Volver a subir el mismo archivo no duplica nada: actualiza las horas "
        "estimadas al valor del archivo.",
    ), start=1):
        ayuda.cell(row=i, column=1, value=linea)
    ayuda.column_dimensions["A"].width = 100

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
